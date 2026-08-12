"""키워드 조사 보고서 DOCX writer."""
from __future__ import annotations

from docx import Document as Docx

from ..models_keyword_report import KeywordReportModel
from .keyword_layout import summary_labels, toc_items
from .keyword_sections import fallback_text


def write_keyword_docx(model: KeywordReportModel, out_path: str, charts: dict | None = None) -> str:
    del charts
    doc = Docx()
    doc.add_paragraph("2026년 여론분석컨설팅 리포트")
    doc.add_paragraph("CONTENTS")
    for item in toc_items():
        doc.add_paragraph(item)

    doc.add_paragraph("REPORT SUMMARY")
    values = [
        model.summary.analysis_period,
        str(model.summary.channel_count),
        model.summary.primary_keyword,
        str(model.summary.positive_ratio),
        model.summary.month_over_month,
    ]
    for label, value in zip(summary_labels(), values):
        doc.add_paragraph(f"{label}: {value}")

    _render_overview(doc, model)
    _render_issues(doc, model)
    _render_channels(doc, model)
    _render_sentiment(doc, model)
    _render_recommendations(doc, model)
    _render_sources(doc, model)

    doc.save(out_path)
    return out_path


def _render_overview(doc: Docx, model: KeywordReportModel) -> None:
    doc.add_paragraph("제1장 분석 개요")
    doc.add_paragraph(f"분석기간 {model.summary.analysis_period}")
    doc.add_paragraph(f"분석채널 {model.summary.channel_count}개")
    doc.add_paragraph(f"핵심 검색어 {model.summary.primary_keyword}")


def _render_issues(doc: Docx, model: KeywordReportModel) -> None:
    doc.add_paragraph("제2장 월간 주요 이슈 동향")
    if not model.issues:
        doc.add_paragraph(fallback_text(False))
        return
    for issue in model.issues:
        doc.add_paragraph(f"{issue.rank}. {issue.title}")
        if issue.subtitle:
            doc.add_paragraph(issue.subtitle)
        if issue.body:
            doc.add_paragraph(issue.body)


def _render_channels(doc: Docx, model: KeywordReportModel) -> None:
    doc.add_paragraph("제3장 온라인 채널별 여론 동향")
    if not model.channel_overview:
        doc.add_paragraph(fallback_text(False))
    else:
        for item in model.channel_overview:
            kws = ", ".join(item.keywords)
            doc.add_paragraph(
                f"{item.channel}: {item.share_or_volume} / {item.change_direction}"
                + (f" / {kws}" if kws else "")
            )
    if model.top_keywords:
        doc.add_paragraph("주요 키워드")
        for item in model.top_keywords:
            doc.add_paragraph(f"{item.rank}. {item.keyword} ({item.category})")


def _render_sentiment(doc: Docx, model: KeywordReportModel) -> None:
    doc.add_paragraph("제4장 감성 분석 - 긍·부정 여론 동향")
    if model.sentiment_overview:
        for label, value in model.sentiment_overview.items():
            doc.add_paragraph(f"{label}: {value}")
    drivers = model.positive_drivers + model.negative_drivers
    if not drivers:
        doc.add_paragraph(fallback_text(False))
        return
    for driver in drivers:
        doc.add_paragraph(driver.title)
        doc.add_paragraph(driver.detail)


def _render_recommendations(doc: Docx, model: KeywordReportModel) -> None:
    doc.add_paragraph("제5장 시사점 및 도정 운영 제언")
    if not model.implications and not model.channel_recommendations:
        doc.add_paragraph(fallback_text(False))
        return
    for item in model.implications:
        doc.add_paragraph(item.get("topic", ""))
        doc.add_paragraph(item.get("headline", ""))
        doc.add_paragraph(item.get("body", ""))
    for item in model.channel_recommendations:
        doc.add_paragraph(f"{item.get('channel', '')}: {item.get('action', '')}")


def _render_sources(doc: Docx, model: KeywordReportModel) -> None:
    doc.add_paragraph("제6장 출처 및 참고자료")
    if model.sources:
        for source in model.sources:
            doc.add_paragraph(f"{source.channel} | {source.title} | {source.url}")
    else:
        doc.add_paragraph(fallback_text(False))
    if model.methodology:
        doc.add_paragraph("방법론")
        for line in model.methodology:
            doc.add_paragraph(line)
