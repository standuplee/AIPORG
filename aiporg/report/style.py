"""DOCX 서식 헬퍼 (공공기관 보고서 톤)."""
from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# 색상
NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x00, 0x70, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)

HEADER_FILL = "1F4E79"     # 표 헤더 음영
BAND_FILL = "DEEAF1"       # 구분 셀 음영
KO_FONT = "맑은 고딕"       # 한국어 기본 글꼴 (Word/한글 공통)


def set_font(run, name: str = KO_FONT, size: float | None = None,
             bold: bool = False, color: RGBColor | None = None):
    run.font.name = name
    run.font.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def shade_cell(cell, fill_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold=False, size=9.5, color=None,
                  align=WD_ALIGN_PARAGRAPH.LEFT, font=KO_FONT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    for i, line in enumerate(str(text).split("\n")):
        if i > 0:
            run = p.add_run()
            run.add_break()
        else:
            run = p.add_run(line)
            set_font(run, font, size, bold, color)
            continue
        run = p.add_run(line)
        set_font(run, font, size, bold, color)


def vcenter_cell(cell):
    from docx.enum.table import WD_ALIGN_VERTICAL
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
