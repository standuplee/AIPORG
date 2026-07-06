"""DOCX 보고서 생성기 — 원본 리포트 절 구성(Ⅰ~Ⅵ) 재현."""
from __future__ import annotations

import os

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt

from ..models import ReportModel
from ..llm.narrator import Narrator
from . import style as S

ROMAN = ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ", "Ⅵ", "Ⅶ", "Ⅷ"]


class DocxBuilder:
    def __init__(self, model: ReportModel, narrator: Narrator, charts: dict,
                 youtube_report=None):
        self.model = model
        self.m = model.meta
        self.a = model.analysis
        self.n = narrator
        self.charts = charts or {}
        self.youtube = youtube_report        # 선택: YouTubeReport
        self.doc = Docx()
        self._base_style()

    # ------------------------------------------------------------------ #
    def _base_style(self):
        st = self.doc.styles["Normal"]
        st.font.name = S.KO_FONT
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(S.qn("w:eastAsia"), S.KO_FONT)
        for sec in self.doc.sections:
            sec.left_margin = Inches(0.9)
            sec.right_margin = Inches(0.9)
            sec.top_margin = Inches(0.9)
            sec.bottom_margin = Inches(0.9)

    # ---- 저수준 헬퍼 -------------------------------------------------- #
    def para(self, text="", *, size=10.5, bold=False, color=S.DARK,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, indent=None, bullet=False):
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        if indent is not None:
            p.paragraph_format.left_indent = Pt(indent)
        prefix = "▪ " if bullet else ""
        run = p.add_run(prefix + text)
        S.set_font(run, S.KO_FONT, size, bold, color)
        return p

    def body_para(self, text: str):
        """개조식 본문 문단 (앞에 ○ 불릿)."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.first_line_indent = Pt(-12)
        run = p.add_run("○ ")
        S.set_font(run, S.KO_FONT, 10.5, False, S.NAVY)
        run2 = p.add_run(text)
        S.set_font(run2, S.KO_FONT, 10.5, False, S.DARK)
        return p

    def section_title(self, idx: int, title: str):
        self.doc.add_page_break()
        # 로마숫자 배지 + 제목
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(f"{ROMAN[idx]}  ")
        S.set_font(r1, S.KO_FONT, 20, True, S.NAVY)
        r2 = p.add_run(title)
        S.set_font(r2, S.KO_FONT, 20, True, S.DARK)
        self._hrule()

    def sub_title(self, num: str, title: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run(f"{num}  ")
        S.set_font(r1, S.KO_FONT, 13, True, S.BLUE)
        r2 = p.add_run(title)
        S.set_font(r2, S.KO_FONT, 13, True, S.DARK)

    def caption(self, text: str):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"[{text}]")
        S.set_font(run, S.KO_FONT, 10, True, S.NAVY)

    def footnote(self, text: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        S.set_font(run, S.KO_FONT, 8.5, False, S.GRAY)

    def image(self, key: str, width=6.2):
        path = self.charts.get(key)
        if path and os.path.exists(path):
            self.doc.add_picture(path, width=Inches(width))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _hrule(self):
        from docx.oxml.ns import qn
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pbdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pbdr.makeelement(qn("w:bottom"), {})
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "18")
        bottom.set(qn("w:color"), "1F4E79")
        pbdr.append(bottom)
        pPr.append(pbdr)

    def table(self, headers, rows, widths=None, band_first_col=False):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = 1
        hdr = t.rows[0].cells
        for i, h in enumerate(headers):
            S.shade_cell(hdr[i], S.HEADER_FILL)
            S.set_cell_text(hdr[i], h, bold=True, size=9.5, color=S.WHITE,
                            align=WD_ALIGN_PARAGRAPH.CENTER)
            S.vcenter_cell(hdr[i])
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                if band_first_col and i == 0:
                    S.shade_cell(cells[i], S.BAND_FILL)
                    S.set_cell_text(cells[i], val, bold=True, size=9.5,
                                    align=WD_ALIGN_PARAGRAPH.CENTER)
                else:
                    S.set_cell_text(cells[i], val, size=9.5, align=align)
                S.vcenter_cell(cells[i])
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        return t

    # ================================================================== #
    # 표지 / 목차
    # ================================================================== #
    def cover(self):
        for _ in range(4):
            self.doc.add_paragraph()
        self.para(f"{self.m.year}년 {self.m.month}월", size=18, bold=True,
                  color=S.NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        self.para("여론분석 컨설팅 리포트", size=30, bold=True, color=S.DARK,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
        self.para(self.m.report_title, size=15, bold=False, color=S.GRAY,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
        for _ in range(6):
            self.doc.add_paragraph()
        self.para(f"{self.m.period_start.strftime('%Y.%m')}", size=13,
                  color=S.NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        self.para(f"분석도구  {self.m.tool_name}", size=11, color=S.GRAY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    def contents(self):
        self.doc.add_page_break()
        self.para("CONTENTS", size=22, bold=True, color=S.NAVY, space_after=14)
        toc = [
            ("Ⅰ", "분석 개요", []),
            ("Ⅱ", f"{self.m.month}월 주요 이슈 동향",
             [f"1. {self.m.month}월 {self.m.subject_name} 온라인 언급량 개요",
              f"2. {self.m.month}월 {self.m.subject_name} 핵심 이슈 TOP10"]),
            ("Ⅲ", "온라인 여론 동향",
             ["1. 채널별 온라인 여론 감성 분석",
              f"2. {self.m.month}월 주요 연관 키워드 및 온라인 특이 동향"]),
            ("Ⅳ", f"{self.m.subject_name} 핵심 키워드 분석", []),
            ("Ⅴ", f"{self.m.leader_title} 직무 수행 평가", []),
            ("Ⅵ", "종합 시사점 및 제언", []),
        ]
        if self.youtube is not None:
            toc.append(("Ⅶ", "유튜브 채널 여론 분석",
                        ["1. 유튜브 채널 언급 개요", "2. 주요 키워드 및 토픽",
                         "3. 핵심 영상", "4. 특이 동향", "5. 직무수행 평가(유튜브 한정)",
                         "6. 유튜브 채널 시사점"]))
        for rn, title, subs in toc:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r1 = p.add_run(f"{rn}   ")
            S.set_font(r1, S.KO_FONT, 13, True, S.BLUE)
            r2 = p.add_run(title)
            S.set_font(r2, S.KO_FONT, 13, True, S.DARK)
            for s in subs:
                sp = self.doc.add_paragraph()
                sp.paragraph_format.left_indent = Pt(28)
                sp.paragraph_format.space_after = Pt(2)
                rr = sp.add_run(s)
                S.set_font(rr, S.KO_FONT, 10.5, False, S.GRAY)

    # ================================================================== #
    # Ⅰ 분석 개요
    # ================================================================== #
    def section_overview(self):
        self.section_title(0, "분석 개요")
        for s in self.n.overview_purpose():
            self.body_para(s)

        self.caption("표 1-1  분석 개요")
        kw_all = self.m.keywords.get(self.m.subject_name, [])
        self.table(
            ["항목", "세부내용"],
            [
                ["분석 방법", "공개 API·RSS 기반 관련 키워드 자동 수집 후 중복제거·"
                             "키워드추출·감성분석·토픽분류·이슈스코어링 수행\n"
                             "키워드 추출(Kiwi·TF-IDF·TextRank) 및 감성분석·LLM 서술 자동화 방법 활용"],
                ["분석 채널", self._channel_desc()],
                ["분석 기간", f"{self.m.period_start.year}년 {self.m.period_start.month}월 "
                             f"{self.m.period_start.day}일 ~ {self.m.period_end.month}월 "
                             f"{self.m.period_end.day}일 {self.m.period_days}일간"],
                ["분석 도구", self._tool_desc()],
            ],
            widths=[1.3, 5.0], band_first_col=True,
        )

        self.caption("표 1-2  키워드 수집 체계")
        rows = [[grp, ", ".join(kws)] for grp, kws in self.m.keywords.items()]
        self.table(["구 분", "내   용"], rows, widths=[1.6, 4.7], band_first_col=True)
        groups = list(self.m.keywords.keys())
        if len(groups) >= 3:
            self.footnote(
                f"* 포괄관계: 「{groups[0]}」이슈는 「{groups[1]}」, 「{groups[2]}」이슈를 "
                f"포함하며, 「{groups[1]}」이슈는 「{groups[2]}」이슈를 포괄합니다.")

    # ================================================================== #
    # Ⅱ 주요 이슈 동향
    # ================================================================== #
    def section_issues(self):
        self.section_title(1, f"{self.m.month}월 주요 이슈 동향")

        self.sub_title("01", f"{self.m.month}월 {self.m.subject_name} 온라인 언급량 개요")
        for s in self.n.volume_overview():
            self.body_para(s)
        self.caption(f"그림 2-1  {self.m.month}월 {self.m.subject_name} 온라인 언급량")
        if self.charts.get("volume_dash"):
            self.image("volume_dash")          # KPI 카드 + 주차별 추이 + 채널 비중 통합
        else:
            self.image("volume")
            self.image("share", width=4.2)
        self.footnote(self._period_tool())

        self.sub_title("02", f"{self.m.month}월 {self.m.subject_name} 핵심 이슈 TOP 10")
        for s in self.n.top_issues_overview():
            self.body_para(s)
        self.caption(f"표 2-1  {self.m.month}월 {self.m.subject_name} 핵심 이슈 TOP 10")
        rows = []
        for i in self.a.top_issues:
            sign = f"+{i.change_pct:.0f}%" if i.change_pct >= 0 else f"{i.change_pct:.0f}%"
            rows.append([str(i.rank), i.category, i.title, f"{i.volume:,}건",
                         sign, i.sentiment.value])
        self.table(["순위", "구분", "이슈", "건수", "증감", "반응"], rows,
                   widths=[0.5, 0.7, 3.2, 0.85, 0.7, 0.6])
        self.footnote(self._period_tool())

    # ================================================================== #
    # Ⅲ 온라인 여론 동향
    # ================================================================== #
    def section_opinion(self):
        self.section_title(2, "온라인 여론 동향")

        self.sub_title("01", "채널별 온라인 여론 감성 분석")
        self.caption("그림 3-1  채널별 감성 비율(%)")
        self.image("sentiment")
        self.footnote(self._period_tool())
        cs = self.n.channel_sentiment()
        for s in cs[:2]:              # 앞 문단(가변 개수 안전 처리)
            self.body_para(s)
        self.caption("그림 3-2  채널별 여론 특징")
        self.image("feature", width=5.5)
        self.footnote(f"* 분석기간 : {self._period()}")
        for s in cs[2:]:             # 나머지 문단
            self.body_para(s)

        self.sub_title("02", f"{self.m.month}월 주요 연관 키워드 및 온라인 특이 동향")
        for s in self.n.related_keywords():
            self.body_para(s)
        self.caption(f"그림 3-3  {self.m.subject_name} 상위 연관 키워드({self.m.month}월)")
        if self.charts.get("keyword_rank"):
            self.image("keyword_rank")         # 랭킹 카드(2단) + 감성 배지
        else:
            self.image("wordcloud")
        self.footnote(f"* 분석기간 : {self._period()}")

        self.caption(f"표 3-1  {self.m.month}월 특이 온라인 동향")
        rows = [[t.title, t.detail] for t in self.a.special_trends]
        self.table(["구   분", "내   용"], rows, widths=[2.1, 4.2], band_first_col=True)
        self.footnote(f"* 분석기간 : {self._period()}")
        for s in self.n.special_trends_summary():
            self.body_para(s)

    # ================================================================== #
    # Ⅳ 핵심 키워드 분석
    # ================================================================== #
    def section_core(self):
        self.section_title(3, f"{self.m.subject_name} 핵심 키워드 분석")
        intro = self.n.core_keywords_intro()
        if not intro:
            self.body_para("핵심 키워드 산출을 위한 언급 표본이 부족함")
            return
        for s in intro:
            self.body_para(s)
        n = len(self.a.core_keywords)
        self.caption(f"그림 4-1  {n}대 핵심 이슈 온라인 여론 상세 분석")
        self.image("core")
        self.footnote(self._period_tool())
        # 카드형(요약·논지 포함) 사용 시 상세 문단은 카드에 담기므로 생략
        has_cards = any(c.summary or c.positive_point for c in self.a.core_keywords)
        if not has_cards:
            for s in self.n.core_keywords_detail():
                self.body_para(s)

    # ================================================================== #
    # Ⅴ 직무수행 평가
    # ================================================================== #
    def section_jobeval(self):
        self.section_title(4, f"{self.m.leader_title} 직무수행 평가")
        lines = self.n.job_evaluation()
        if not lines:
            self.body_para(f"{self.m.leader_name} {self.m.leader_title} 관련 언급 표본이 "
                           f"부족하여 직무수행 평가는 산출하지 않음 "
                           f"(수집 채널 확대 또는 분석기간 조정 시 산출 가능)")
            return
        self.caption(f"그림 5-1  {self._leader_full()} 직무수행 평가"
                     f"({self.m.year}년 {self.m.month}월)")
        self.image("jobeval", width=4.5)
        self.footnote(self._period_tool())
        for s in lines:
            self.body_para(s)

    # ================================================================== #
    # Ⅵ 종합 시사점 및 제언
    # ================================================================== #
    def section_implications(self):
        self.section_title(5, "종합 시사점 및 제언")
        if not self.a.implications:
            self.body_para("종합 시사점 및 대응 제언은 수집·집계 결과에 대한 전문 분석가 "
                           "검토 또는 LLM 서술 엔진(gemini/openai) 연동 단계에서 생성됨")
            self.body_para("본 실데이터 시연은 자동 수집·집계·감성분석 결과까지를 포함하며, "
                           "정책 제언 등 해석 산출물은 별도 단계로 분리 운영함")
            # 참고: 자동 도출된 핵심 이슈 요약 제공
            if self.a.top_issues:
                self.caption("참고  자동 도출 핵심 이슈 요약")
                rows = [[str(i.rank), i.category, i.title, f"{i.volume:,}건", i.sentiment.value]
                        for i in self.a.top_issues[:5]]
                self.table(["순위", "구분", "이슈", "언급량", "감성"], rows,
                           widths=[0.5, 0.7, 3.6, 0.9, 0.6])
            return
        self.caption(f"표 6-1  {self.m.subject_name} 종합 시사점 및 대응 제언"
                     f"({self.m.year}년 {self.m.month}월)")
        rows = []
        for idx, imp in enumerate(self.a.implications, 1):
            body = f"{imp.body}\n\n▷ 대응 : {imp.response}"
            rows.append([f"{idx:02d}", f"{imp.topic}\n\n[{imp.headline}]", body])
        self.table(["No.", "시사점", "대응 제언"], rows,
                   widths=[0.5, 1.9, 4.0], band_first_col=True)
        for s in self.n.implications_conclusion():
            self.body_para(s)

    # ------------------------------------------------------------------ #
    CHANNEL_DESC = {
        "뉴스기사": "뉴스기사(네이버·구글)",
        "블로그": "블로그(네이버)",
        "커뮤니티": "커뮤니티(디시인사이드·클리앙·에펨코리아·더쿠·루리웹)",
        "유튜브": "유튜브(YouTube Data API)",
        "공공RSS": "공공 RSS·정부 보도자료",
        "X(트위터)": "X(트위터)",
        "인스타그램": "인스타그램",
    }

    def _tool_desc(self) -> str:
        """분석 도구를 실제 코드 구성에 맞게 기술."""
        collectors = "Naver 검색 API(뉴스·블로그)·Google News·공공 RSS·커뮤니티 어댑터"
        if self.youtube is not None:
            collectors += "·YouTube Data API v3"
        return (f"AI 여론분석 플랫폼 {self.m.tool_name} — 수집({collectors}), "
                f"분석(Kiwi·TF-IDF·TextRank 키워드 추출, 감성분석, 토픽 군집, 이슈 스코어링), "
                f"보고서 서술(LLM) 활용")

    def _channel_desc(self) -> str:
        """실제 수집된(=channel_share 에 존재하는) 채널만 기재 + 제외 채널 명시."""
        keys = list(self.a.channel_share.keys()) or ["뉴스기사"]
        desc = ", ".join(self.CHANNEL_DESC.get(k, k) for k in keys)
        excluded = [c for c in ("X(트위터)", "인스타그램") if c not in keys]
        if excluded:
            desc += ("\n※ " + "·".join(excluded) +
                     "은 공식 API 정책상 대량 수집이 제한되어 분석 대상에서 제외")
        return desc

    def _leader_full(self) -> str:
        """지자체명 + 직위 결합 시 '도'/'시' 중복 제거. 예: 충청남도+도지사->충청남도지사."""
        s, t = self.m.subject_name, self.m.leader_title
        if s and t and s[-1] == t[0]:
            return s + t[1:]
        return s + t

    def _period(self) -> str:
        s, e = self.m.period_start, self.m.period_end
        return f"{s.year}년 {s.month}월 {s.day}일 ~ {e.month}월 {e.day}일"

    def _period_tool(self) -> str:
        return f"* 분석기간 : {self._period()}  ⦀  수집도구 : {self.m.tool_name}"

    # ================================================================== #
    # Ⅶ 유튜브 채널 여론 분석 (선택 — YouTubeReport 제공 시)
    # ================================================================== #
    def section_youtube(self):
        r = self.youtube
        self.section_title(6, "유튜브 채널 여론 분석")
        self.body_para("본 장은 온라인 여론의 하나의 채널인 유튜브에 한정한 분석으로, "
                       "전체 여론을 대표하지 않으며 유튜브 이용자 중심의 반응으로 해석함")

        # 01 개요
        self.sub_title("01", "유튜브 채널 언급 개요")
        ov = r.youtube_overview
        self.body_para(ov.narrative)
        self.caption("표 7-1  유튜브 수집·감성 개요")
        cs = r.channel_sentiment
        self.table(
            ["항목", "내용"],
            [["수집 영상", f"{ov.video_count:,}건"],
             ["수집 댓글", f"{ov.comment_count:,}건"],
             ["누적 조회수", f"{ov.total_views:,}회"],
             ["댓글 감성", f"긍정 {cs.positive}% / 중립 {cs.neutral}% / "
                          f"부정 {cs.negative}% / 혼재 {cs.mixed}%"]],
            widths=[1.5, 4.8], band_first_col=True)
        self.caption("그림 7-1  유튜브 영상 댓글 감성 분포")
        self.image("yt_sentiment", width=4.2)
        self.footnote(self._period_tool())

        # 02 키워드·토픽
        self.sub_title("02", "유튜브 주요 키워드 및 토픽")
        kws = ", ".join(k.keyword for k in r.top_keywords[:8])
        self.body_para(f"유튜브 채널 내에서는 {kws} 등의 키워드가 주요하게 나타남")
        self.image("yt_keywords", width=5.6)
        self.image("yt_topics", width=5.6)
        self.footnote(self._period_tool())

        # 03 핵심 영상
        self.sub_title("03", "유튜브 핵심 영상")
        self.caption("표 7-2  조회수 상위 영상")
        rows = []
        for i, v in enumerate(r.top_videos[:8], 1):
            rows.append([str(i), v.channel_title[:14], v.title[:40],
                         f"{v.view_count:,}", v.sentiment or "중립"])
        self.table(["순위", "채널", "제목", "조회수", "감성"], rows,
                   widths=[0.5, 1.2, 3.3, 0.9, 0.6])
        self.footnote(self._period_tool())

        # 04 특이 동향
        if r.notable_trends:
            self.sub_title("04", "유튜브 특이 동향")
            self.caption("표 7-3  유튜브 채널 특이 동향")
            self.table(["구   분", "내   용"],
                       [[t.title, t.detail] for t in r.notable_trends],
                       widths=[2.0, 4.3], band_first_col=True)

        # 05 도지사 평가(유튜브 한정)
        self.sub_title("05", f"{self.m.leader_title} 직무수행 평가 (유튜브 한정)")
        ge = r.governor_evaluation
        self.body_para(f"{ge.subject}에 대한 유튜브 채널 내 반응은 긍정 {ge.positive_pct}%, "
                       f"부정 {ge.negative_pct}%, 중립 {ge.neutral_pct}%로 집계됨")
        for p in ge.positive_points:
            self.body_para(f"(긍정) {p}")
        for p in ge.negative_points:
            self.body_para(f"(부정) {p}")
        self.footnote("※ " + ge.note)

        # 06 시사점 및 AI 요약
        self.sub_title("06", "유튜브 채널 시사점")
        for rec in r.recommendations:
            self.body_para(rec)
        if r.ai_summary:
            self.caption("AI 요약")
            self.body_para(r.ai_summary)

    # ------------------------------------------------------------------ #
    def build(self, out_path: str) -> str:
        self.cover()
        self.contents()
        self.section_overview()
        self.section_issues()
        self.section_opinion()
        self.section_core()
        self.section_jobeval()
        self.section_implications()
        if self.youtube is not None:
            self.section_youtube()
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        self.doc.save(out_path)
        return out_path


def write_docx(model: ReportModel, narrator: Narrator, charts: dict, out_path: str,
               youtube_report=None) -> str:
    return DocxBuilder(model, narrator, charts, youtube_report).build(out_path)
